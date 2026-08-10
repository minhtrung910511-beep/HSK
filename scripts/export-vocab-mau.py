"""Export vocab-data.ts sang Excel theo mẫu: STT | Hán tự | Phiên âm | Nghĩa | Chủ đề"""
import re, os, sys
sys.path.insert(0, "/home/z/my-project/skills/xlsx")
sys.path.insert(0, "/home/z/my-project/skills/xlsx/templates")

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Đọc vocab-data.ts
with open("/home/z/my-project/src/lib/vocab-data.ts", "r", encoding="utf-8") as f:
    content = f.read()

# Parse TOPICS để map topic_id → topic_name
topic_pattern = re.compile(r'\{\s*id:\s*"([^"]+)",\s*name:\s*"([^"]+)"')
topic_map = {}
for m in topic_pattern.finditer(content[:content.index("export const VOCAB")]):
    topic_map[m.group(1)] = m.group(2)

# Parse VOCAB entries
vocab_pattern = re.compile(
    r'\{\s*id:\s*(\d+),\s*han:\s*"([^"]+)",\s*pinyin:\s*"([^"]+)",\s*meaning:\s*"([^"]+)",\s*pos:\s*"([^"]+)",\s*emoji:\s*"[^"]*"(?:,\s*imageUrl:\s*"[^"]*")?,\s*example:\s*"([^"]+)",\s*examplePinyin:\s*"([^"]+)",\s*exampleVi:\s*"([^"]+)",\s*topic:\s*"([^"]+)"',
    re.DOTALL
)

words = []
for m in vocab_pattern.finditer(content):
    topic_name = topic_map.get(m.group(9), m.group(9))
    words.append({
        "id": int(m.group(1)),
        "han": m.group(2),
        "pinyin": m.group(3),
        "meaning": m.group(4),
        "pos": m.group(5),
        "topic": topic_name,
        "topic_id": m.group(9),
    })

# Sắp xếp theo topic rồi theo id
words.sort(key=lambda w: (w["topic_id"], w["id"]))

print(f"Total words: {len(words)}")
print(f"Topics: {len(topic_map)}")

# ===== Tạo Excel =====
wb = Workbook()
ws = wb.active
ws.title = "Từ vựng HSK1"

# Styles - theo mẫu
header_font = Font(name="Times New Roman", size=12, bold=True, color="FF0000", italic=True)
header_fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

data_font = Font(name="Times New Roman", size=11, italic=True)
han_font = Font(name="Times New Roman", size=14, bold=True, color="CC0000")
stt_font = Font(name="Times New Roman", size=11, bold=True)
data_fill = PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid")
center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

thin_border = Border(
    left=Side(style="thin", color="388E3C"),
    right=Side(style="thin", color="388E3C"),
    top=Side(style="thin", color="388E3C"),
    bottom=Side(style="thin", color="388E3C"),
)

# ===== Header =====
headers = ["STT", "Hán tự", "Phiên âm", "Nghĩa", "Chủ đề"]
for col_idx, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border

# ===== Data rows =====
for row_idx, w in enumerate(words, 2):
    row_data = [row_idx - 1, w["han"], w["pinyin"], w["meaning"], w["topic"]]
    
    for col_idx, value in enumerate(row_data, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.fill = data_fill
        cell.border = thin_border
        
        if col_idx == 1:  # STT
            cell.font = stt_font
            cell.alignment = center_align
        elif col_idx == 2:  # Hán tự
            cell.font = han_font
            cell.alignment = center_align
        elif col_idx == 3:  # Phiên âm
            cell.font = Font(name="Times New Roman", size=11, italic=True)
            cell.alignment = left_align
        elif col_idx == 4:  # Nghĩa
            cell.font = Font(name="Times New Roman", size=11, italic=True)
            cell.alignment = left_align
        elif col_idx == 5:  # Chủ đề
            cell.font = Font(name="Times New Roman", size=11, italic=True)
            cell.alignment = left_align

# ===== Column widths =====
ws.column_dimensions["A"].width = 6    # STT
ws.column_dimensions["B"].width = 14   # Hán tự
ws.column_dimensions["C"].width = 20   # Phiên âm
ws.column_dimensions["D"].width = 30   # Nghĩa
ws.column_dimensions["E"].width = 22   # Chủ đề

# ===== Row heights =====
ws.row_dimensions[1].height = 30
for row_idx in range(2, len(words) + 2):
    ws.row_dimensions[row_idx].height = 28

# ===== Freeze + Auto filter =====
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:E{len(words) + 1}"

# ===== Workbook metadata =====
wb.properties.creator = "Z.ai"
wb.properties.title = "Từ vựng HSK1 - 426 từ"
wb.properties.description = f"{len(words)} từ vựng HSK1 theo {len(topic_map)} chủ đề"

# ===== Save =====
output_path = "/home/z/my-project/download/tu-vung-hsk1-mau.xlsx"
wb.save(output_path)
print(f"\n✅ Đã lưu: {output_path}")
print(f"   {len(words)} dòng × 5 cột")
print(f"   File size: {os.path.getsize(output_path) / 1024:.1f} KB")
